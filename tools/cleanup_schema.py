#!/usr/bin/env python3
"""Synchronize database schema with current import spreadsheets.

Reads the three base Excel files in ``imports/`` and ensures the tables
``produtos``, ``fichas_tecnicas`` and ``precos_taxas`` match their headers.
A backup of the SQLite database (``*.db.bak``) is created before applying
changes.
"""

from pathlib import Path
from shutil import copy2

from openpyxl import load_workbook

from utils.paths import get_project_root
from data.datastore import DataStore
from data.migration import setup_database
from services.products import canonicalize_header, sync_table_schema


def main() -> None:
    root = get_project_root()
    imports_dir = root / "imports"

    files = {
        "produtos": imports_dir / "Produtos_Base.xlsx",
        "fichas_tecnicas": imports_dir / "FichasTecnicas_base.xlsx",
        "precos_taxas": imports_dir / "PreçosTaxas_base.xlsx",
    }

    ds = DataStore()
    conn = getattr(ds, "conn", None)
    if conn is None:
        return

    db_file = Path(conn.execute("PRAGMA database_list").fetchone()[2])
    backup = db_file.with_suffix(db_file.suffix + ".bak")
    copy2(db_file, backup)

    setup_database(conn)

    for table, path in files.items():
        if not path.exists():
            continue
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = ws.iter_rows(values_only=True)
        try:
            raw_headers = list(next(rows))
        except StopIteration:
            wb.close()
            continue
        if table == "precos_taxas":
            mapped = [canonicalize_header(h, table=table) for h in raw_headers]
            existing = [
                r[1].lower() for r in conn.execute(f"PRAGMA table_info({table})")
            ]
            headers = mapped + [c for c in existing if c not in mapped]
        else:
            headers = raw_headers
        sync_table_schema(conn, table, headers)
        wb.close()

    conn.commit()
    ds.close()


if __name__ == "__main__":
    main()
