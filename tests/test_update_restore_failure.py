from __future__ import annotations

import sqlite3
from pathlib import Path
from unittest.mock import patch

import pytest
from openpyxl import Workbook

from services.products import ProductService


def _write_update_files(base_dir: Path) -> None:
    produtos = base_dir / "Produtos_Base.xlsx"
    fichas = base_dir / "FichasTecnicas_base.xlsx"
    precos = base_dir / "PreçosTaxas_base.xlsx"

    wb_prod = Workbook()
    ws = wb_prod.active
    ws.append(["Codigo", "Produto", "TipoVenda"])
    ws.append(["P1", "Produto 1 atualizado", 1])
    wb_prod.save(produtos)
    wb_prod.close()

    wb_ft = Workbook()
    ws_ft = wb_ft.active
    ws_ft.append(
        [
            "FamiliaSubfamilia",
            "ProdutoCodigo",
            "ProdutoNome",
            "ComponenteCodigo",
            "ComponenteNome",
            "Qtd",
            "Unidade",
            "Ppu",
            "Preco",
            "Peso",
            "Ordem",
        ]
    )
    ws_ft.append(
        [
            None,
            "P1",
            "Produto 1 atualizado",
            "I1",
            "Ingrediente 1",
            1,
            "Un",
            2,
            2,
            1,
            1,
        ]
    )
    wb_ft.save(fichas)
    wb_ft.close()

    wb_prec = Workbook()
    ws_prec = wb_prec.active
    ws_prec.append(["Codigo", "Preco1", "Preco2", "Preco3", "Preco4", "Preco5"])
    ws_prec.append(["P1", 9, None, None, None, None])
    wb_prec.save(precos)
    wb_prec.close()


def _snapshot_tables(db_path: Path, tables: list[str]) -> dict[str, list[tuple]]:
    snapshot: dict[str, list[tuple]] = {}
    with sqlite3.connect(db_path) as conn:
        conn.row_factory = sqlite3.Row
        for table in tables:
            rows = conn.execute(f"SELECT * FROM {table}").fetchall()
            normalized: list[tuple] = []
            for row in rows:
                values = []
                for value in row:
                    if isinstance(value, memoryview):
                        values.append(bytes(value))
                    else:
                        values.append(value)
                normalized.append(tuple(values))
            snapshot[table] = sorted(normalized)
    return snapshot


class SimpleDataStore:
    def __init__(self, db_path: Path):
        self.db_path = Path(db_path)
        self.conn = sqlite3.connect(self.db_path)
        self.conn.row_factory = sqlite3.Row
        self._ids: list[str] = []
        self.reload_ids()

    def reload_ids(self) -> None:
        if self.conn is None:
            self._ids = []
            return
        cur = self.conn.execute("SELECT Codigo FROM Produtos ORDER BY Codigo")
        self._ids = [row[0] for row in cur.fetchall()]

    def close(self) -> None:
        if self.conn is not None:
            self.conn.close()
            self.conn = None


def _prepare_database(db_path: Path) -> None:
    conn = sqlite3.connect(db_path)
    conn.execute(
        "CREATE TABLE IF NOT EXISTS Produtos ("
        "Codigo TEXT PRIMARY KEY, Produto TEXT, TipoVenda INTEGER, Preco1G REAL,"
        " Preco2G REAL, Iva REAL)"
    )
    conn.execute(
        "CREATE TABLE IF NOT EXISTS FichasTecnicas ("
        "ProdutoCodigo TEXT, ProdutoNome TEXT, ComponenteCodigo TEXT,"
        " ComponenteNome TEXT, Qtd REAL, Unidade TEXT, Ppu REAL,"
        " Preco REAL, Peso REAL, Ordem INTEGER)"
    )
    conn.execute(
        "CREATE TABLE IF NOT EXISTS PrecosTaxas ("
        "Codigo TEXT, Loja TEXT, Preco1 REAL, Preco2 REAL, Preco3 REAL,"
        " Preco4 REAL, Preco5 REAL)"
    )
    conn.execute(
        "CREATE TABLE IF NOT EXISTS Uploads ("
        "Id INTEGER PRIMARY KEY AUTOINCREMENT, Filename TEXT, Content BLOB)"
    )
    conn.commit()
    conn.close()


def test_update_from_excel_restores_backup_on_failure(tmp_path, monkeypatch):
    monkeypatch.setattr("services.products.get_project_root", lambda: tmp_path)
    monkeypatch.setattr("data.backup.get_project_root", lambda: tmp_path)
    monkeypatch.setattr("services.products.setup_database", lambda conn: None)

    imports_dir = tmp_path / "imports"
    imports_dir.mkdir(parents=True, exist_ok=True)
    logs_dir = tmp_path / "logs"
    logs_dir.mkdir(parents=True, exist_ok=True)
    databases_dir = tmp_path / "databases"
    databases_dir.mkdir(parents=True, exist_ok=True)
    db_path = databases_dir / "ftv.db"

    _prepare_database(db_path)
    ds = SimpleDataStore(db_path)
    service = ProductService(ds)

    conn = ds.conn
    conn.execute(
        "INSERT INTO Produtos (Codigo, Produto, TipoVenda) VALUES (?, ?, 1)",
        ("P1", "Produto 1"),
    )
    conn.execute(
        "INSERT INTO FichasTecnicas (ProdutoCodigo, ProdutoNome, ComponenteCodigo, "
        "ComponenteNome, Qtd, Unidade, Ppu, Preco, Peso, Ordem) "
        "VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        ("P1", "Produto 1", "I1", "Ingrediente 1", 1, "Un", 2, 2, 1, 1),
    )
    conn.execute(
        "INSERT INTO PrecosTaxas (Codigo, Loja, Preco1) VALUES (?, ?, ?)",
        ("P1", "L1", 5),
    )
    conn.commit()
    ds.reload_ids()

    tables = ["Produtos", "FichasTecnicas", "PrecosTaxas", "Uploads"]
    before = _snapshot_tables(db_path, tables)

    _write_update_files(imports_dir)

    with patch(
        "openpyxl.workbook.workbook.Workbook.save", side_effect=RuntimeError("boom")
    ):
        with pytest.raises(RuntimeError, match="boom"):
            service.update_from_excel()

    after = _snapshot_tables(db_path, tables)
    assert after == before

    assert isinstance(service.ds, SimpleDataStore)
    assert service.ds.conn is not None
    with sqlite3.connect(db_path) as check_conn:
        row = check_conn.execute(
            "SELECT Produto FROM Produtos WHERE Codigo='P1'"
        ).fetchone()
    assert row[0] == "Produto 1"
