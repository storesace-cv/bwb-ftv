import shutil
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from data.datastore import DataStore
import services.products as products
from services.products import ProductService, _update_from_excel, _import_single_excel
from utils.paths import get_project_root

PRODUCTS_FILE = products.IMPORT_FILE_BASENAMES["Produtos"]
FICHAS_FILE = products.IMPORT_FILE_BASENAMES["FichasTecnicas"]
PRECOS_FILE = products.IMPORT_FILE_BASENAMES["PrecosTaxas"]


@pytest.fixture
def ds():
    ds = DataStore(":memory:")
    conn = ds.conn
    conn.execute("DROP TABLE Produtos")
    conn.execute(
        "CREATE TABLE Produtos (Codigo TEXT PRIMARY KEY, Produto TEXT, "
        "TipoVenda INTEGER, Preco1G REAL, Preco2G REAL, Iva REAL)"
    )
    conn.execute("DROP TABLE FichasTecnicas")
    conn.execute(
        "CREATE TABLE FichasTecnicas ("
        "FamiliaSubfamilia TEXT, ProdutoCodigo TEXT, ProdutoNome TEXT, "
        "ComponenteCodigo TEXT, ComponenteNome TEXT, Qtd REAL, Unidade TEXT, "
        "Ppu REAL, Preco REAL, Peso REAL, Ordem INTEGER)"
    )
    return ds


@pytest.fixture
def imports_dir():
    base = get_project_root() / "imports"
    if base.exists():
        shutil.rmtree(base)
    base.mkdir(parents=True, exist_ok=True)
    yield base
    shutil.rmtree(base, ignore_errors=True)


@pytest.fixture
def logs_dir():
    base = get_project_root() / "logs"
    if base.exists():
        shutil.rmtree(base)
    yield base
    shutil.rmtree(base, ignore_errors=True)


def _write_base_files(
    base_dir, products=None, code_header="Codigo", prices=None
):
    if products is None:
        products = [("P1", "Produto 1")]
    prod_wb = Workbook()
    ws = prod_wb.active
    ws.append(["Codigo", "Produto", "tipo_venda"])
    for code, name in products:
        ws.append([code, name, 1])
    prod_wb.save(base_dir / PRODUCTS_FILE)

    ft_wb = Workbook()
    ws = ft_wb.active
    ws.append(
        [
            "familia_subfamilia",
            "produto_codigo",
            "produto_nome",
            "componente_codigo",
            "componente_nome",
            "Qtd",
            "Unidade",
            "Ppu",
            "Preco",
            "Peso",
            "Ordem",
        ]
    )
    for code, _ in products:
        ws.append([None, code, None, None, None, None, None, None, None, None, None])
    ft_wb.save(base_dir / FICHAS_FILE)

    if prices is None:
        prices = [1.0, None, None, None, None]
    prec_wb = Workbook()
    ws = prec_wb.active
    ws.append([code_header, "Preco1", "Preco2", "Preco3", "Preco4", "Preco5"])
    ws.append([products[0][0], *prices])
    prec_wb.save(base_dir / PRECOS_FILE)


def _write_sparse_codigo_files(base_dir, products):
    prod_wb = Workbook()
    ws = prod_wb.active
    ws.append(["Codigo", "Produto", "tipo_venda"])
    for code, name, _ in products:
        ws.append([code, name, 1])
    prod_wb.save(base_dir / PRODUCTS_FILE)

    ft_wb = Workbook()
    ws = ft_wb.active
    ws.append(
        [
            "familia_subfamilia",
            "produto_codigo",
            "produto_nome",
            "componente_codigo",
            "componente_nome",
            "Qtd",
            "Unidade",
            "Ppu",
            "Preco",
            "Peso",
            "Ordem",
        ]
    )
    for code, name, components in products:
        for ordem, (comp_code, comp_name) in enumerate(components, start=1):
            ws.append(
                [
                    None,
                    code if ordem == 1 else None,
                    name if ordem == 1 else None,
                    comp_code,
                    comp_name,
                    ordem,
                    "Un",
                    None,
                    None,
                    None,
                    ordem,
                ]
            )
    ft_wb.save(base_dir / FICHAS_FILE)

    prec_wb = Workbook()
    ws = prec_wb.active
    ws.append(["Codigo", "Preco1", "Preco2", "Preco3", "Preco4", "Preco5"])
    for code, _, _ in products:
        ws.append([code, 1, None, None, None, None])
    prec_wb.save(base_dir / PRECOS_FILE)


def test_import_from_excel_replaces_database(ds, imports_dir):
    ds.conn.execute("INSERT INTO Produtos (Codigo, Produto) VALUES ('OLD', 'Old')")
    ds.reload_ids()
    svc = ProductService(ds)
    _write_base_files(imports_dir, products=[("P1", "Produto 1"), ("P2", "Produto 2")])
    svc.import_from_excel()
    ds.reload_ids()

    assert ds.total() == 2
    assert ds.get_produto_info("OLD") == {}
    assert ds.get_produto_info("P1")["produto"] == "Produto 1"
    assert ds.get_produto_info("P2")["produto"] == "Produto 2"
    assert not (imports_dir / PRODUCTS_FILE).exists()
    rows = ds.conn.execute("SELECT Filename FROM Uploads").fetchall()
    names = {r[0] for r in rows}
    assert PRODUCTS_FILE in names


def test_update_from_excel_updates_and_inserts(ds, imports_dir):
    ds.conn.executemany(
        "INSERT INTO Produtos (Codigo, Produto, TipoVenda) VALUES (?, ?, 1)",
        [("P1", "Produto 1"), ("P2", "Produto 2")],
    )
    ds.conn.executemany(
        "INSERT INTO FichasTecnicas (ProdutoCodigo) VALUES (?)",
        [("P1",), ("P2",)],
    )
    ds.reload_ids()
    svc = ProductService(ds)
    _write_base_files(
        imports_dir,
        products=[("P1", "Produto 1 updated"), ("P3", "Produto 3")],
    )
    svc.update_from_excel()
    ds.reload_ids()

    assert ds.total() == 3
    assert ds.get_produto_info("P1")["produto"] == "Produto 1 updated"
    assert ds.get_produto_info("P2")["produto"] == "Produto 2"
    assert ds.get_produto_info("P3")["produto"] == "Produto 3"
    assert not (imports_dir / PRODUCTS_FILE).exists()
    rows = ds.conn.execute("SELECT Filename FROM Uploads").fetchall()
    names = {r[0] for r in rows}
    assert PRODUCTS_FILE in names


def test_import_from_excel_missing_file(ds, imports_dir):
    _write_base_files(imports_dir)
    (imports_dir / PRECOS_FILE).unlink()
    svc = ProductService(ds)
    with pytest.raises(FileNotFoundError) as exc:
        svc.import_from_excel()
    assert PRECOS_FILE in str(exc.value)


def test_update_from_excel_missing_file(ds, imports_dir):
    _write_base_files(imports_dir)
    (imports_dir / FICHAS_FILE).unlink()
    svc = ProductService(ds)
    with pytest.raises(FileNotFoundError) as exc:
        svc.update_from_excel()
    assert FICHAS_FILE in str(exc.value)


def test_import_stores_files_in_uploads(ds, imports_dir):
    svc = ProductService(ds)
    _write_base_files(imports_dir)
    svc.import_from_excel()
    rows = ds.conn.execute("SELECT Filename, length(Content) FROM Uploads").fetchall()
    assert len(rows) == 3
    expected = [PRODUCTS_FILE, FICHAS_FILE, PRECOS_FILE]
    names = {r[0] for r in rows}
    for name in expected:
        assert name in names
        assert not (imports_dir / name).exists()


def test_update_stores_files_in_uploads(ds, imports_dir):
    svc = ProductService(ds)
    _write_base_files(imports_dir)
    svc.update_from_excel()
    rows = ds.conn.execute("SELECT Filename, length(Content) FROM Uploads").fetchall()
    assert len(rows) == 3
    expected = [PRODUCTS_FILE, FICHAS_FILE, PRECOS_FILE]
    names = {r[0] for r in rows}
    for name in expected:
        assert name in names
        assert not (imports_dir / name).exists()


def test_update_from_excel_attempts_backup_before_update(monkeypatch, ds):
    svc = ProductService(ds)
    monkeypatch.setattr(ProductService, "_should_create_backup", lambda self: True)
    calls: list[tuple[str, str | None]] = []

    def fake_backup(*, prefix=None) -> Path:
        calls.append(("backup", prefix))
        return Path("dummy-backup.db")

    def fake_update(datastore) -> Path:
        assert datastore is ds
        calls.append(("update", None))
        return Path("dummy-report.xlsx")

    monkeypatch.setattr(products, "create_backup", fake_backup)
    monkeypatch.setattr(products, "update_from_excel", fake_update)

    result = svc.update_from_excel()

    assert result == Path("dummy-report.xlsx")
    assert calls == [("backup", "ftv-actualizacao-"), ("update", None)]


def test_update_from_excel_aborts_when_backup_fails(monkeypatch, ds):
    svc = ProductService(ds)
    monkeypatch.setattr(ProductService, "_should_create_backup", lambda self: True)
    called = {"update": False}

    def fake_backup(*, prefix=None) -> Path:
        assert prefix == "ftv-actualizacao-"
        raise RuntimeError("backup falhou")

    def fake_update(datastore):
        called["update"] = True

    monkeypatch.setattr(products, "create_backup", fake_backup)
    monkeypatch.setattr(products, "update_from_excel", fake_update)

    with pytest.raises(RuntimeError):
        svc.update_from_excel()

    assert called["update"] is False


def test_update_from_excel_generates_report(ds, imports_dir, logs_dir):
    ds.conn.execute(
        "INSERT INTO Produtos (Codigo, Produto, TipoVenda, Preco1G) "
        "VALUES ('P1', 'Produto 1', 1, 10)"
    )
    ds.conn.execute(
        "INSERT INTO FichasTecnicas "
        "(ProdutoCodigo, ComponenteCodigo, ComponenteNome, Qtd, Unidade, Ppu, Preco, Ordem) "
        "VALUES ('P1', 'I1', 'Ingrediente 1', 1, 'Kg', 2, 2, 1)"
    )
    ds.conn.execute(
        "INSERT INTO PrecosTaxas (Codigo, Preco1, Loja) VALUES ('P1', 10, '1')"
    )
    ds.reload_ids()

    prod_wb = Workbook()
    ws = prod_wb.active
    ws.append(["Codigo", "Produto", "TipoVenda", "Preco1G"])
    ws.append(["P1", "Produto 1 Atualizado", 1, 15])
    ws.append(["P2", "Produto 2", 1, 20])
    prod_wb.save(imports_dir / PRODUCTS_FILE)

    ft_wb = Workbook()
    ws = ft_wb.active
    ws.append(
        [
            "familia_subfamilia",
            "produto_codigo",
            "produto_nome",
            "componente_codigo",
            "componente_nome",
            "Qtd",
            "Unidade",
            "Ppu",
            "Preco",
            "Peso",
            "Ordem",
        ]
    )
    ws.append([None, "P1", "Produto 1 Atualizado", "I1", "Ingrediente 1", 2, "Kg", 3, 6, None, 1])
    ws.append([None, None, None, "I2", "Ingrediente 2", 1, "Un", None, None, None, 2])
    ft_wb.save(imports_dir / FICHAS_FILE)

    prec_wb = Workbook()
    ws = prec_wb.active
    ws.append(["Codigo", "Preco1", "Preco2", "Preco3", "Preco4", "Preco5"])
    ws.append(["P1", 12, None, None, None, None])
    ws.append(["P2", 22, None, None, None, None])
    prec_wb.save(imports_dir / PRECOS_FILE)

    svc = ProductService(ds)
    report_path = svc.update_from_excel()

    assert report_path is not None
    assert report_path.exists()
    assert report_path.parent == logs_dir

    wb = load_workbook(report_path)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()

    assert rows[0] == (
        "Tipo",
        "Tabela",
        "Produto",
        "Componente",
        "Nome Componente",
        "Alterações",
    )

    def _find(tipo, tabela, produto, componente=None):
        for row in rows[1:]:
            if row[0] == tipo and row[1] == tabela and row[2] == produto:
                if componente is None or row[3] == componente:
                    return row
        return None

    produto_update = _find("Atualização de Registo", "Produtos", "P1")
    assert produto_update is not None
    assert "Produto 1 -> Produto 1 Atualizado" in (produto_update[5] or "")

    novo_produto = _find("Novo Registo", "Produtos", "P2")
    assert novo_produto is not None
    assert "Codigo: P2" in (novo_produto[5] or "")

    ingrediente_update = _find("Atualização de Ingrediente", "FichasTecnicas", "P1", "I1")
    assert ingrediente_update is not None
    assert "Qtd: 1 -> 2" in (ingrediente_update[5] or "")

    novo_ingrediente = _find("Novo Ingrediente", "FichasTecnicas", "P1", "I2")
    assert novo_ingrediente is not None
    assert "ComponenteNome: Ingrediente 2" in (novo_ingrediente[5] or "")

    preco_update = _find("Atualização de Registo", "PrecosTaxas", "P1")
    assert preco_update is not None
    assert "Preco1: 10 -> 12" in (preco_update[5] or "")

    novo_preco = _find("Novo Registo", "PrecosTaxas", "P2")
    assert novo_preco is not None
    assert "Codigo: P2" in (novo_preco[5] or "")


def test_import_from_excel_uses_produto_codigo(ds, imports_dir):
    _write_base_files(
        imports_dir, code_header="produto_codigo", prices=[2.5, None, None, None, None]
    )
    svc = ProductService(ds)
    svc.import_from_excel()
    pvps = ds.get_pvps("P1")
    assert pvps["pvps"][0] == 2.5
    assert len(pvps["pvps"]) == 5


def test_update_from_excel_uses_produto_codigo(ds, imports_dir):
    ds.conn.execute(
        "INSERT INTO Produtos (Codigo, Produto, Preco1G) VALUES ('P1', 'X', 1.0)"
    )
    ds.reload_ids()
    _write_base_files(
        imports_dir, code_header="produto_codigo", prices=[3.0, None, None, None, None]
    )
    svc = ProductService(ds)
    svc.update_from_excel()
    pvps = ds.get_pvps("P1")
    assert pvps["pvps"][0] == 3.0
    assert len(pvps["pvps"]) == 5


def test_import_from_excel_skips_rows_without_codigo(ds, imports_dir):
    products = [("P1", "Produto 1"), ("   ", "Sem Código"), (None, "Outro")]
    _write_base_files(imports_dir, products=products)
    prec_path = imports_dir / PRECOS_FILE
    wb = load_workbook(prec_path)
    ws = wb.active
    ws.append([" ", None, None, None, None, None])
    wb.save(prec_path)
    wb.close()

    svc = ProductService(ds)
    svc.import_from_excel()

    rows = ds.conn.execute(
        "SELECT Codigo FROM Produtos WHERE Codigo IS NULL OR Codigo = ''"
    ).fetchall()
    assert rows == []

    ft_rows = ds.conn.execute(
        "SELECT ProdutoCodigo FROM FichasTecnicas "
        "WHERE ProdutoCodigo IS NULL OR ProdutoCodigo = ''"
    ).fetchall()
    assert ft_rows == []


def test_update_from_excel_skips_rows_without_codigo(ds, imports_dir):
    ds.conn.execute(
        "INSERT INTO Produtos (Codigo, Produto, TipoVenda) VALUES ('P1', 'Produto 1', 1)"
    )
    ds.conn.execute(
        "INSERT INTO FichasTecnicas (ProdutoCodigo, ProdutoNome) VALUES ('P1', 'Produto 1')"
    )
    ds.reload_ids()

    products = [("P1", "Produto 1 atualizado"), ("P2", "Produto 2"), ("", "Sem Código")]
    _write_base_files(imports_dir, products=products)
    prec_path = imports_dir / PRECOS_FILE
    wb = load_workbook(prec_path)
    ws = wb.active
    ws.append(["", None, None, None, None, None])
    wb.save(prec_path)
    wb.close()

    svc = ProductService(ds)
    svc.update_from_excel()
    ds.reload_ids()

    codigos = {
        row[0]
        for row in ds.conn.execute("SELECT Codigo FROM Produtos").fetchall()
        if row[0] is not None
    }
    assert "" not in codigos
    assert "P1" in codigos and "P2" in codigos

    preco_rows = ds.conn.execute(
        "SELECT Codigo FROM PrecosTaxas WHERE Codigo IS NULL OR Codigo = ''"
    ).fetchall()
    assert preco_rows == []

    ft_invalid = ds.conn.execute(
        "SELECT ProdutoCodigo FROM FichasTecnicas "
        "WHERE ProdutoCodigo IS NULL OR ProdutoCodigo = ''"
    ).fetchall()
    assert ft_invalid == []


def test_import_from_excel_propagates_blank_produto_codigo(ds, imports_dir):
    products = [
        ("P1", "Produto 1", [("I1", "Ingrediente 1"), ("I2", "Ingrediente 2")]),
        ("P2", "Produto 2", [("I3", "Ingrediente 3"), ("I4", "Ingrediente 4")]),
    ]
    _write_sparse_codigo_files(imports_dir, products)

    svc = ProductService(ds)
    svc.import_from_excel()
    ds.reload_ids()

    for code, _, components in products:
        ingredientes = ds.get_ingredientes(code)
        assert [i["ComponenteNome"] for i in ingredientes] == [
            comp_name for _, comp_name in components
        ]


def test_update_from_excel_propagates_blank_produto_codigo(ds, imports_dir):
    ds.conn.execute(
        "INSERT INTO Produtos (Codigo, Produto, TipoVenda) VALUES ('P1', 'Antigo', 1)"
    )
    ds.conn.execute(
        "INSERT INTO FichasTecnicas (ProdutoCodigo, ComponenteNome, Ordem) "
        "VALUES ('P1', 'Velho', 1)"
    )
    ds.reload_ids()

    products = [
        ("P1", "Produto 1", [("I1", "Ingrediente 1"), ("I2", "Ingrediente 2")]),
        ("P2", "Produto 2", [("I3", "Ingrediente 3"), ("I4", "Ingrediente 4")]),
    ]
    _write_sparse_codigo_files(imports_dir, products)

    svc = ProductService(ds)
    svc.update_from_excel()
    ds.reload_ids()

    for code, _, components in products:
        ingredientes = ds.get_ingredientes(code)
        nomes = [i["ComponenteNome"] for i in ingredientes]
        assert nomes == [comp_name for _, comp_name in components]
        assert "Velho" not in nomes

def test_import_single_excel_accepts_canonical_headers(ds, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Codigo", "Produto", "Preco1G", "Preco2G", "Iva"])
    ws.append(["P1", "Produto 1", 10.5, None, 23])
    ws.append(["P2", "Produto 2", 5.75, 7.25, None])
    path = tmp_path / "produtos.xlsx"
    wb.save(path)

    _import_single_excel(path, ds)

    rows = [
        tuple(row)
        for row in ds.conn.execute(
            "SELECT Codigo, Produto, Preco1G, Preco2G, Iva FROM Produtos ORDER BY Codigo"
        ).fetchall()
    ]

    assert rows == [
        ("P1", "Produto 1", 10.5, None, 23),
        ("P2", "Produto 2", 5.75, 7.25, None),
    ]


def test_import_from_excel_reads_all_prices(ds, imports_dir):
    _write_base_files(imports_dir, prices=[1, 2, 3, 4, 5])
    svc = ProductService(ds)
    svc.import_from_excel()
    pvps = ds.get_pvps("P1")
    assert pvps["pvps"] == [1.0, 2.0, 3.0, 4.0, 5.0]


def test_update_from_excel_reads_all_prices(ds, imports_dir):
    ds.conn.execute(
        "INSERT INTO Produtos (Codigo, Produto, Preco1G) VALUES ('P1', 'X', 1.0)"
    )
    ds.reload_ids()
    _write_base_files(imports_dir, prices=[6, 7, 8, 9, 10])
    svc = ProductService(ds)
    svc.update_from_excel()
    pvps = ds.get_pvps("P1")
    assert pvps["pvps"] == [6.0, 7.0, 8.0, 9.0, 10.0]


@pytest.mark.parametrize(
    "price,expected",
    [("1,5", 1.5), ("1 234,5", 1234.5)],
)
def test_import_from_excel_parses_formatted_numbers(ds, imports_dir, price, expected):
    _write_base_files(imports_dir, prices=[price] * 5)
    svc = ProductService(ds)
    svc.import_from_excel()
    pvps = ds.get_pvps("P1")
    assert pvps["pvps"] == [expected] * 5


@pytest.mark.parametrize(
    "price,expected",
    [("1,5", 1.5), ("1 234,5", 1234.5)],
)
def test_update_from_excel_parses_formatted_numbers(ds, imports_dir, price, expected):
    ds.conn.execute(
        "INSERT INTO Produtos (Codigo, Produto, Preco1G) VALUES ('P1', 'X', 1.0)"
    )
    ds.reload_ids()
    _write_base_files(imports_dir, prices=[price] * 5)
    svc = ProductService(ds)
    svc.update_from_excel()
    pvps = ds.get_pvps("P1")
    assert pvps["pvps"] == [expected] * 5


@pytest.mark.parametrize("func", ["import_from_excel", "update_from_excel"])
def test_preco_taxas_requires_codigo(ds, imports_dir, func):
    _write_base_files(imports_dir, code_header="wrong")
    svc = ProductService(ds)
    with pytest.raises(ValueError):
        getattr(svc, func)()


def test_import_reads_preco(ds, imports_dir):
    prod = Workbook()
    ws = prod.active
    ws.append(["codigo", "nome", "tipo_venda"])
    ws.append(["P1", "Produto 1", 1])
    prod.save(imports_dir / PRODUCTS_FILE)

    ft = Workbook()
    ws = ft.active
    ws.append(
        [
            "familia_subfamilia",
            "produto_codigo",
            "produto_nome",
            "componente_codigo",
            "componente_nome",
            "Qtd",
            "Unidade",
            "Ppu",
            "Preco",
            "Peso",
            "Ordem",
        ]
    )
    ws.append([None, "P1", None, None, "Ing", 2, "Kg", 3, "1 234,5", None, None])
    ft.save(imports_dir / FICHAS_FILE)

    prec = Workbook()
    ws = prec.active
    ws.append(["codigo", "Preco1", "Preco2", "Preco3", "Preco4", "Preco5"])
    ws.append(["P1", 1, None, None, None, None])
    prec.save(imports_dir / PRECOS_FILE)

    svc = ProductService(ds)
    svc.import_from_excel()
    ing = ds.get_ingredientes("P1")[0]
    assert ing["Preco"] == 1234.5


def test_update_reads_preco(ds, imports_dir):
    ds.conn.execute("INSERT INTO Produtos (Codigo, Produto) VALUES ('P1', 'Prod')")
    ds.conn.execute(
        "INSERT INTO FichasTecnicas "
        "(ProdutoCodigo, ComponenteNome, Qtd, Unidade, Ppu, Preco) "
        "VALUES ('P1', 'Ing', 1, 'Kg', 2, 5)"
    )
    ds.reload_ids()

    prod = Workbook()
    ws = prod.active
    ws.append(["codigo", "nome", "tipo_venda"])
    ws.append(["P1", "Prod", 1])
    prod.save(imports_dir / PRODUCTS_FILE)

    ft = Workbook()
    ws = ft.active
    ws.append(
        [
            "familia_subfamilia",
            "produto_codigo",
            "produto_nome",
            "componente_codigo",
            "componente_nome",
            "Qtd",
            "Unidade",
            "Ppu",
            "Preco",
            "Peso",
            "Ordem",
        ]
    )
    ws.append([None, "P1", None, None, "Ing", 3, "Kg", 2, 7, None, None])
    ft.save(imports_dir / FICHAS_FILE)

    prec = Workbook()
    ws = prec.active
    ws.append(["codigo", "Preco1", "Preco2", "Preco3", "Preco4", "Preco5"])
    ws.append(["P1", 1, None, None, None, None])
    prec.save(imports_dir / PRECOS_FILE)

    svc = ProductService(ds)
    svc.update_from_excel()
    ing = ds.get_ingredientes("P1")[0]
    assert ing["Preco"] == 7


def test_private_update_from_excel_parses_numbers(ds, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Código", "Produto", "Preço1 G"])
    ws.append(["P1", "Produto 1", "1 234,5"])
    path = tmp_path / "update.xlsx"
    wb.save(path)

    _update_from_excel(path, ds)
    info = ds.get_produto_info("P1")
    assert info["produto"] == "Produto 1"
    assert info["preco1g"] == 1234.5

