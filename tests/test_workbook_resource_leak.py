"""Test that workbook resources are properly closed even when exceptions occur."""

import os
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook


def test_workbook_closes_on_exception_in_try_finally():
    """Verify that workbooks are closed even when exceptions occur in try-finally blocks."""
    
    # Create a temporary Excel file
    import tempfile
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = Path(tmp.name)
    
    try:
        # Create a test workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Header1", "Header2"])
        ws.append(["Value1", "Value2"])
        wb.save(tmp_path)
        wb.close()
        
        # Simulate the pattern used in services/products.py
        wb = load_workbook(tmp_path, read_only=True, data_only=True)
        try:
            ws = wb.active
            rows = ws.iter_rows(values_only=True)
            headers = list(next(rows))
            
            # Simulate an exception during processing
            raise RuntimeError("Simulated error during processing")
        finally:
            wb.close()
            
    except RuntimeError:
        # Expected exception
        pass
    
    # Verify the file can be deleted (not locked by an open workbook)
    # On Windows, this would fail if the file handle wasn't properly closed
    tmp_path.unlink()
    assert not tmp_path.exists()


def test_workbook_closes_on_early_return_in_try_finally():
    """Verify that workbooks are closed on early returns in try-finally blocks."""
    
    # Create a temporary Excel file
    import tempfile
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = Path(tmp.name)
    
    try:
        # Create an empty workbook (will trigger StopIteration)
        wb = Workbook()
        wb.save(tmp_path)
        wb.close()
        
        # Simulate the pattern used in services/products.py
        def process_workbook(file_path):
            wb = load_workbook(file_path, read_only=True, data_only=True)
            try:
                ws = wb.active
                rows = ws.iter_rows(values_only=True)
                try:
                    headers = list(next(rows))
                except StopIteration:
                    # Early return - workbook should still be closed
                    return
                # More processing would happen here
            finally:
                wb.close()
        
        process_workbook(tmp_path)
        
    finally:
        # Verify the file can be deleted (not locked)
        tmp_path.unlink()
        assert not tmp_path.exists()


def test_workbook_resource_leak_without_finally():
    """Demonstrate that without try-finally, workbooks can leak on exceptions."""
    
    # Create a temporary Excel file
    import tempfile
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tmp:
        tmp_path = Path(tmp.name)
    
    try:
        # Create a test workbook
        wb = Workbook()
        ws = wb.active
        ws.append(["Header1", "Header2"])
        ws.append(["Value1", "Value2"])
        wb.save(tmp_path)
        wb.close()
        
        # Simulate the OLD pattern (without try-finally)
        def process_without_finally(file_path):
            wb = load_workbook(file_path, read_only=True, data_only=True)
            ws = wb.active
            rows = ws.iter_rows(values_only=True)
            headers = list(next(rows))
            
            # If an exception occurs here, wb.close() is never called
            raise RuntimeError("Simulated error")
            
            # This line would never be reached
            wb.close()  # pragma: no cover
        
        try:
            process_without_finally(tmp_path)
        except RuntimeError:
            pass
        
        # The workbook is still open! On some systems, we might not be able to delete it
        # But we can at least verify the pattern is problematic
        # (In practice, Python's garbage collector may eventually close it, but that's not guaranteed)
        
    finally:
        # Clean up - may need to force close
        try:
            tmp_path.unlink()
        except PermissionError:
            # On Windows, the file might still be locked
            # This demonstrates the resource leak issue
            import gc
            gc.collect()  # Force garbage collection
            try:
                tmp_path.unlink()
            except PermissionError:
                # Still locked - this is the bug we're fixing
                pass
